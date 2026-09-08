#!/usr/bin/env python3
"""Generate warehouse group/item seeds from Warehouses.xlsx."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / "src/app/mocks/data/seed"
XLSX = Path("/Users/omarahmedamin/Downloads/Warehouses.xlsx")
CHUNK = 80
SKIP_NAMES = {"الصنف", "الوحدة", "الصتف", None, ""}

GROUPS = [
    ("oilSeal", "اويل سيل", "Oil Seal"),
    ("plumbing", "سباكة", "Plumbing"),
    ("belts", "سيور", "Belts"),
    ("screws", "مسامير+فولة غاطسة", "Screws & Countersunk Washers"),
    ("ledgers", "دفاتر", "Stationery"),
    ("oils", "زيوت + شحوم", "Oils & Grease"),
    ("bearings", "رولمان بلى", "Ball Bearings"),
    ("mfgMaterials", "خامات للتصنيع", "Manufacturing Materials"),
    ("campaign", "الحملة", "Campaign"),
    ("safety", "الامن الصناعي", "Industrial Safety"),
    ("paints", "دهانات", "Paints"),
    ("electrical", "الكهرباء", "Electrical"),
    ("mechanical", "الميكانيكا", "Mechanical"),
    ("air", "الهواء", "Air"),
    ("hall", "الصالة", "Production Floor"),
    ("tools", "عدد وادوات", "Tools & Equipment"),
]

SUBS = [
    ("electrical", "contactor", "كونتاكتور", "Contactor"),
    ("electrical", "fuse", "فيوز", "Fuse"),
    ("electrical", "switches", "مفاتيح", "Switches"),
    ("electrical", "lamps", "لمبات+كشافات", "Lamps & Floodlights"),
    ("electrical", "cables", "كابلات", "Cables"),
    ("electrical", "devices", "اجهزة كهرياء+حساسات", "Electrical Devices & Sensors"),
    ("electrical", "relay", "ريلاي+اوفر لود", "Relay & Overload"),
    ("electrical", "motors", "مواتير", "Motors"),
    ("electrical", "terminals", "كوس+ترامل", "Terminals"),
    ("electrical", "elecMisc", "متنوع كلاسك", "Electrical Miscellaneous"),
    ("mechanical", "coupling", "كوبلن", "Coupling"),
    ("mechanical", "rodBar", "تيش + بار", "Rod & Bar"),
    ("mechanical", "gaskets", "جوانات", "Gaskets"),
    ("mechanical", "weldWire", "سلك لحام", "Welding Wire"),
    ("mechanical", "flange", "فلانشة", "Flange"),
    ("mechanical", "cooler", "كولر", "Cooler"),
    ("mechanical", "weldElbow", "كوع لحام", "Welding Elbow"),
    ("mechanical", "weldReducer", "مسلوب لحام", "Welding Reducer"),
    ("mechanical", "packing", "حشو", "Packing"),
    ("mechanical", "flexAir", "وصلات مرنة+قرب هواء", "Flexible Connections & Air Tanks"),
    ("mechanical", "mechSeal", "ميكانيكل سيل+تيل", "Mechanical Seal & Pins"),
    ("mechanical", "pipes", "مواسير", "Pipes"),
    ("mechanical", "ironware", "حدايد", "Ironware"),
    ("mechanical", "lathe", "مستلزمات مخرطة", "Lathe Supplies"),
    ("mechanical", "pumps", "طلمبات", "Pumps"),
    ("mechanical", "hydraulic", "مستلزمات هيدرولك", "Hydraulic Supplies"),
    ("mechanical", "valves", "محابس", "Valves"),
    ("mechanical", "mechMisc", "متنوع ميكانيكا", "Mechanical Miscellaneous"),
    ("air", "gauges", "عدادات", "Gauges"),
    ("air", "airConn", "وصلات هواء", "Air Connections"),
    ("air", "airDevices", "اجهزة هواء", "Air Devices"),
    ("air", "compressor", "كمبروسر", "Compressor"),
    ("hall", "quality", "الجودة", "Quality"),
    ("hall", "production", "الانتاج", "Production"),
    ("hall", "prep", "التحضيرات", "Preparations"),
    ("tools", "handTools", "عدد", "Tools"),
    ("tools", "equipment", "ادوات", "Equipment"),
    ("tools", "bits", "بنط", "Drill Bits"),
]

GROUP_ALIASES = {
    "اويل سيل": "oilSeal",
    "اوبل سيل": "oilSeal",
    "سباكة": "plumbing",
    "سيور": "belts",
    "السيور": "belts",
    "مسامير+فولة غاطسة": "screws",
    "المسامير + فولة غاطسة": "screws",
    "دفاتر": "ledgers",
    "دفاتروادوات مكتبية": "ledgers",
    "زيوت + شحوم": "oils",
    "زيوت وشحوم": "oils",
    "رولمان بلى": "bearings",
    "رولمان بلي": "bearings",
    "خامات للتصنيع": "mfgMaterials",
    "الحملة": "campaign",
    "الامن الصناعي": "safety",
    "دهانات": "paints",
    "الكهرباء": "electrical",
    "الميكانيكا": "mechanical",
    "الهواء": "air",
    "الصالة": "hall",
    "عدد وادوات": "tools",
}

SUB_ALIASES = {
    "كونتاكتور": "contactor",
    "فيوز": "fuse",
    "مفاتيح": "switches",
    "لمبات+كشافات": "lamps",
    "لمبات + كشافات": "lamps",
    "كابلات": "cables",
    "اجهزة كهرياء+حساسات": "devices",
    "اجهزة + حساسات": "devices",
    "اجهزة كهرباء + حساسات": "devices",
    "ريلاي+اوفر لود": "relay",
    "ريلاي + اوفر لود": "relay",
    "ريلاى+اوفر لود": "relay",
    "مواتير": "motors",
    "كوس+ترامل": "terminals",
    "كوس + ترامل": "terminals",
    "متنوع كلاسك": "elecMisc",
    "متنوع كلاسيك": "elecMisc",
    "متنوع كهرباء كلاسيك": "elecMisc",
    "كوبلن": "coupling",
    "تيش + بار": "rodBar",
    "جوانات": "gaskets",
    "جوانات+اورنج": "gaskets",
    "سلك لحام": "weldWire",
    "فلانشة": "flange",
    "كولر": "cooler",
    "كوع لحام": "weldElbow",
    "مسلوب لحام": "weldReducer",
    "حشو": "packing",
    "وصلات مرنة+قرب هواء": "flexAir",
    "وصلات مرنة + قرب هواء": "flexAir",
    "ميكانيكل سيل+تيل": "mechSeal",
    "ميكانيكل سيل + تيل": "mechSeal",
    "تيل+ميكانيكل سيل": "mechSeal",
    "مواسير": "pipes",
    "حدايد": "ironware",
    "مستلزمات مخرطة": "lathe",
    "طلمبات": "pumps",
    "مستلزمات هيدرولك": "hydraulic",
    "مستلزمات هيدروليك": "hydraulic",
    "محابس": "valves",
    "متنوع ميكانيكا": "mechMisc",
    "عدادات": "gauges",
    "وصلات هواء": "airConn",
    "اجهزة هواء": "airDevices",
    "كمبروسر": "compressor",
    "كومبروسر": "compressor",
    "الجودة": "quality",
    "الجودة+المقص": "quality",
    "الانتاج": "production",
    "التحضيرات": "prep",
    "عدد": "handTools",
    "ادوات": "equipment",
    "أدوات": "equipment",
    "بنط": "bits",
}

UNITS = {
    "عدد": "units.count",
    "قطعة": "units.piece",
    "ك/ع": "units.box",
    "لتر": "units.liter",
    "كيلو": "units.kg",
    "كجم": "units.kg",
    "سم": "units.cm",
    "متر": "units.m",
    "باكو": "units.pack",
    "طن": "units.ton",
    "لفة": "units.roll",
}


def compact(value: str) -> str:
    return re.sub(r"[\s+]+", "", str(value).strip())


def lookup(name: object, table: dict[str, str], kind: str) -> str:
    text = str(name).strip()
    if text in table:
        return table[text]
    key = compact(text)
    for raw, slug in table.items():
        if compact(raw) == key:
            return slug
    raise SystemExit(f"Unknown {kind} header: {text!r}")


def ts(value: str) -> str:
    return value.replace("\\", "\\\\").replace("'", "\\'")


def cell(ws, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    return "" if value is None else str(value).strip()


def unit_key(raw: str) -> str:
    text = raw.strip()
    if text in UNITS:
        return UNITS[text]
    raise SystemExit(f"Unknown unit: {text!r}")


def parse_spare(ws) -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    parent = ""
    for col in range(1, ws.max_column + 1, 2):
        title = cell(ws, 1, col)
        sub = cell(ws, 2, col)
        if title:
            parent = lookup(title, GROUP_ALIASES, "group")
        if not parent:
            continue
        group = f"itemGroups.{parent}"
        child = ""
        if sub and sub not in SKIP_NAMES:
            child = f"itemSubGroups.{lookup(sub, SUB_ALIASES, 'sub-group')}"
        for row in range(3, ws.max_row + 1):
            name = cell(ws, row, col)
            unit = cell(ws, row, col + 1)
            if name in SKIP_NAMES or not name:
                continue
            rows.append((name, group, child, unit_key(unit or "عدد")))
    return rows


def parse_simple(ws) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for row in range(2, ws.max_row + 1):
        name = cell(ws, row, 1)
        unit = cell(ws, row, 2)
        if name in SKIP_NAMES or not name:
            continue
        rows.append((name, unit_key(unit or "عدد")))
    return rows


def write_chunk(path: Path, export: str, lines: list[str]) -> None:
    body = "\n".join(
        [
            "/** Generated by tools/gen-warehouse-seed.py — do not edit. */",
            "import { StockItem } from '../../../core/models/warehouse.models';",
            "import { stock } from './warehouse-items.util';",
            "",
            f"export const {export}: StockItem[] = [",
            *lines,
            "];",
            "",
        ]
    )
    path.write_text(body, encoding="utf-8")
    count = body.count("\n")
    if count > 300:
        raise SystemExit(f"{path.name} has {count} lines")


def item_line(
    code: str,
    name: str,
    warehouse: str,
    group: str,
    sub: str,
    unit: str,
) -> str:
    latin = re.sub(r"[^\x00-\x7F]+", " ", name).strip()
    english = latin if re.search(r"[A-Za-z]", latin) else name
    return (
        f"  stock('{code}', '{ts(name)}', '{ts(english)}', "
        f"'{warehouse}', '{group}', '{sub}', '{unit}'),"
    )


def write_lookups() -> None:
    lines = [
        "/** Generated by tools/gen-warehouse-seed.py — do not edit. */",
        "import { LookupValue } from '../../../core/models/system.models';",
        "import { TRANSLATIONS } from '../i18n';",
        "",
        "const lk = (",
        "  id: string,",
        "  group: string,",
        "  value: string,",
        "  labelAr: string,",
        "  labelEn: string,",
        "  parentValue?: string,",
        "): LookupValue => {",
        "  TRANSLATIONS['ar'][value] = labelAr;",
        "  TRANSLATIONS['en'][value] = labelEn;",
        "  return { id, group, value, labelAr, labelEn, parentValue };",
        "};",
        "",
        "export const SEED_ITEM_GROUPS: LookupValue[] = [",
    ]
    for slug, ar, en in GROUPS:
        lines.append(
            f"  lk('lv-ig-{slug}', 'itemGroups', 'itemGroups.{slug}', '{ts(ar)}', '{ts(en)}'),"
        )
    lines += ["];", "", "export const SEED_ITEM_SUBGROUPS: LookupValue[] = ["]
    for parent, slug, ar, en in SUBS:
        lines.append(
            f"  lk('lv-is-{slug}', 'itemSubGroups', 'itemSubGroups.{slug}', "
            f"'{ts(ar)}', '{ts(en)}', 'itemGroups.{parent}'),"
        )
    lines += [
        "];",
        "",
        "export const SEED_STOCK_UNITS: LookupValue[] = [",
        "  lk('lv-un-count', 'units', 'units.count', 'عدد', 'Count'),",
        "  lk('lv-un-piece', 'units', 'units.piece', 'قطعة', 'Piece'),",
        "  lk('lv-un-box', 'units', 'units.box', 'ك/ع', 'Carton'),",
        "  lk('lv-un-liter', 'units', 'units.liter', 'لتر', 'Liter'),",
        "  lk('lv-un-kg', 'units', 'units.kg', 'كجم', 'Kilogram'),",
        "  lk('lv-un-ton', 'units', 'units.ton', 'طن', 'Ton'),",
        "  lk('lv-un-cm', 'units', 'units.cm', 'سم', 'Centimetre'),",
        "  lk('lv-un-m', 'units', 'units.m', 'متر', 'Metre'),",
        "  lk('lv-un-pack', 'units', 'units.pack', 'باكو', 'Pack'),",
        "  lk('lv-un-roll', 'units', 'units.roll', 'لفة', 'Roll'),",
        "];",
        "",
    ]
    path = SEED / "item-groups.seed.ts"
    path.write_text("\n".join(lines), encoding="utf-8")
    if path.read_text(encoding="utf-8").count("\n") > 300:
        raise SystemExit(f"{path.name} exceeds 300 lines")


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    spare = parse_spare(wb["مخزن قطع غيار كامل"])
    raw = parse_simple(wb["مخزن الخامات"])
    supplies = parse_simple(wb["مخزن مستلزمات"])
    wb.close()

    for old in SEED.glob("spare-items-*.ts"):
        old.unlink()
    if (SEED / "spare-items.ts").exists():
        (SEED / "spare-items.ts").unlink()

    spare_exports: list[str] = []
    for index, start in enumerate(range(0, len(spare), CHUNK), start=1):
        export = f"SEED_SPARE_{index:02d}"
        spare_exports.append(export)
        lines = [
            item_line(f"SPR-{1000 + start + offset}", name, "wh-spare", group, sub, unit)
            for offset, (name, group, sub, unit) in enumerate(spare[start : start + CHUNK])
        ]
        write_chunk(SEED / f"spare-items-{index:02d}.ts", export, lines)

    barrel = [
        "/** Generated by tools/gen-warehouse-seed.py — do not edit. */",
        "import { StockItem } from '../../../core/models/warehouse.models';",
        *[
            f"import {{ {name} }} from './spare-items-{i:02d}';"
            for i, name in enumerate(spare_exports, start=1)
        ],
        "",
        "export const SEED_SPARE_ITEMS: StockItem[] = [",
        *[f"  ...{name}," for name in spare_exports],
        "];",
        "",
    ]
    (SEED / "spare-items.ts").write_text("\n".join(barrel), encoding="utf-8")

    write_chunk(
        SEED / "raw-items.seed.ts",
        "SEED_RAW_ITEMS",
        [
            item_line(f"RAW-{i:03d}", name, "wh-raw", "", "", unit)
            for i, (name, unit) in enumerate(raw, start=1)
        ],
    )
    write_chunk(
        SEED / "supply-items.seed.ts",
        "SEED_SUPPLY_ITEMS",
        [
            item_line(f"SPL-{i:03d}", name, "wh-supplies", "", "", unit)
            for i, (name, unit) in enumerate(supplies, start=1)
        ],
    )
    write_lookups()

    groups = Counter(row[1] for row in spare)
    subs = Counter(row[2] for row in spare if row[2])
    units = Counter(row[3] for row in spare) + Counter(row[1] for row in raw) + Counter(
        row[1] for row in supplies
    )
    print(f"spare={len(spare)} raw={len(raw)} supplies={len(supplies)}")
    print("groups", dict(groups))
    print("subs", dict(subs))
    print("units", dict(units))


if __name__ == "__main__":
    main()
