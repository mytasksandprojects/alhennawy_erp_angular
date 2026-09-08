#!/usr/bin/env python3
"""Generate AR/EN mock seed files from the Al Hennawy Excel workbooks."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / "src/app/mocks/data/seed"
DOWNLOADS = Path("/Users/omarahmedamin/Downloads")

PHRASES = [
    ("Shareholders' Current Account", "جاري المساهمين"),
    ("Selling and Marketing Expenses", "مصروفات البيع والتسويق"),
    ("Accrued General & Administrative Expenses", "مصروفات إدارية مستحقة"),
    ("Accrued Operating Expenses", "مصروفات تشغيل مستحقة"),
    ("Administrative Expenses", "المصروفات الإدارية"),
    ("Accounts Receivable", "العملاء"),
    ("Accounts Payable", "الموردون"),
    ("Suppliers / Accounts Payable", "الموردون / حسابات دائنة"),
    ("Current Liabilities", "الخصوم المتداولة"),
    ("Current Assets", "الأصول المتداولة"),
    ("Fixed Assets", "الأصول الثابتة"),
    ("Operating Expenses", "مصروفات التشغيل"),
    ("Deferred Revenue", "إيرادات مؤجلة"),
    ("Accrued Expenses", "مصروفات مستحقة"),
    ("Employee Advances", "سلف العاملين"),
    ("Employee Custody", "عهد العاملين"),
    ("Notes Receivable Account", "أوراق القبض"),
    ("Notes Payable", "أوراق الدفع"),
    ("Bank Accounts", "حسابات البنوك"),
    ("Cash Accounts", "حسابات الكاش"),
    ("National Bank of Egypt", "البنك الأهلي المصري"),
    ("Chemical Factory", "مصنع الكيماويات"),
    ("Net Sales", "صافي المبيعات"),
    ("Gross Sales", "إجمالي المبيعات"),
    ("Sales Returns", "مردودات المبيعات"),
    ("Purchase Returns", "مردودات المشتريات"),
    ("Scrap Purchases", "مشتريات دشت"),
    ("Spare Parts Purchases", "مشتريات قطع غيار"),
    ("Chemicals Purchases", "مشتريات كيماويات"),
    ("Oils and Lubricants Purchases", "مشتريات زيوت وشحوم"),
    ("General Purchases", "مشتريات عامة"),
    ("Scrap Purchase Returns", "مردودات مشتريات دشت"),
    ("Finished Products", "منتج تام"),
    ("Local Finished Products", "منتج تام محلي"),
    ("Export Finished Products", "منتج تام تصدير"),
    ("Waste Paper", "ورق دشت"),
    ("Production Supplies", "مستلزمات إنتاج"),
    ("Main Safe", "الخزينة الرئيسية"),
    ("Electrical Appliances", "أجهزة كهربائية"),
    ("Laboratory Parts", "قطع معملية"),
    ("Industrial Security and Professional Safety", "الأمن الصناعي والسلامة المهنية"),
    ("Transportation and Daily Allowances Expenses", "مصروفات انتقال وبدلات يومية"),
    ("Sales Scales", "موازين المبيعات"),
    ("QNB Bank", "بنك قطر الوطني"),
    ("Warehouse – ", "مخزن – "),
    ("Customer – ", "عميل – "),
    ("Supplier – ", "مورد – "),
    ("Advance – ", "سلفة – "),
    ("Custody of ", "عهدة "),
    ("Custody – ", "عهدة – "),
    ("Warehouses", "المخازن"),
    ("Expenses", "المصروفات"),
    ("Revenue", "الإيرادات"),
    ("Liability", "خصوم"),
    ("Liabilities", "الخصوم"),
    ("Equity", "حقوق الملكية"),
    ("Capital", "رأس المال"),
    ("Purchases", "المشتريات"),
    ("Assets", "الأصول"),
    ("Asset", "الأصول"),
]

WORDS = {
    "Lands": "أراضي",
    "Cameras": "كاميرات",
    "Cars": "سيارات",
    "Machines": "آلات",
    "Equipment": "معدات",
    "Buildings": "مباني",
    "furniture": "أثاث",
    "Furniture": "أثاث",
    "Bank": "بنك",
    "Cash": "كاش",
    "Factory": "مصنع",
    "Line": "خط",
    "Driver": "سائق",
    "Warehouse": "مخزن",
    "Maintenance": "صيانة",
    "Electrical": "كهرباء",
    "Security": "أمن",
    "Vehicle": "سيارة",
    "Expenses": "مصروفات",
    "Account": "حساب",
    "Accounts": "حسابات",
    "General": "عام",
    "Operating": "تشغيل",
    "Administrative": "إدارية",
    "Marketing": "تسويق",
    "Selling": "بيع",
    "and": "و",
    "of": "",
    "the": "",
    "for": "لـ",
    "Scrap": "دشت",
    "Company": "شركة",
}

CUS_AR = {
    "Al-Alameya for Tissue Paper": ("العالمية للورق الصحي", "برج العرب الجديدة"),
    "Misk Trading and Supplies Establishment": ("مؤسسة مسك للتجارة والتوريدات", "الخانكة"),
    "Al Rabie Company for Paper Products": ("شركة الربيع لمنتجات الورق", "السادس من أكتوبر"),
    "Al Karma Janet Morris Anton": ("الكرمة جانيت موريس أنطون", "القاهرة"),
    "Al Sharq Company for Tissue Paper Packaging": ("شركة الشرق لتعبئة الورق الصحي", "القليوبية"),
    "Al nawras al fanni for manufacturing company": ("شركة النورس الفني للتصنيع", "جدة"),
    "Al Maaly Tissue Company for paper products": ("شركة المعالي للورق الصحي", "جدة"),
    "TOP PLASTIC FACTORY": ("مصنع توب بلاستيك", "جدة"),
    "Durra AlSalihiya Company For Trading": ("شركة درة الصالحية للتجارة", "جدة"),
    "KHM for gen tradig fabo ow by k h mk": ("كي إتش إم للتجارة العامة", "بيروت"),
    "Al Dawlia for Manufacturing, Cutting and Packaging Paper Products": (
        "الدولية لتصنيع وتقطيع وتعبئة منتجات الورق",
        "الخانكة",
    ),
    "Arabian Company for Industry and Trade(fay)": ("الشركة العربية للصناعة والتجارة (فاي)", "بدر"),
    "Qimma Al Sahab for Paper Manufacturing": ("قمة السحاب لصناعة الورق", "خميس مشيط"),
    "Golden Eagle for Paper Products": ("النسر الذهبي لمنتجات الورق", "القليوبية"),
    "alwaraq alssihhi Factory For Paper Products Establishment": (
        "مؤسسة مصنع الورق الصحي لمنتجات الورق",
        "تبوك",
    ),
    "Tripack Group": ("مجموعة تريباك", "البحيرة"),
}

SUP_EN = {
    "شركة اليسر للزيوت و الشحوم": "Al Yosr Oils & Grease",
    "شركة باور لوب للزيوت و الشحوم": "Power Lube Oils & Grease",
    "شركة فوكس ايجيبت": "Fuchs Egypt",
    "داتكو للخدمات الصناعية": "Datco Industrial Services",
    "الشركة العربية للتوريدات الكهربائية و الميكانيكية": "Arabian Electrical & Mechanical Supplies",
    "الشرق الاوسط للحلول الصناعية": "Middle East Industrial Solutions",
    "ماستر الكتريك": "Master Electric",
    "المصرية للمقاييس": "Egyptian Scales Co.",
    "الوتين لتجارة الاطارات": "Al Wateen Tyres Trading",
    "ال غنيم للمعدات الصناعية": "Al Ghoneim Industrial Equipment",
    "شركة الاتحاد": "Al Ittihad Company",
    "الفا تك": "Alpha Tech",
    "السعد اليكتريك للتوريدات الكهربائية": "Al Saad Electric Supplies",
    "الدولية للسيور": "International Belts",
    "المؤسسة العلمية للهندسة الصناعية": "Scientific Foundation for Industrial Engineering",
    "الصفوة": "Al Safwa",
    "شركة الهندسية للهيدروليك": "Engineering Hydraulics Co.",
    "مصر لصناعة الكيماويات": "Misr Chemical Industries",
    "شركة العزيزية": "Al Aziziyah Company",
    "مؤسسة الشرق للتجارة و التوريدات": "Al Sharq Trading & Supplies",
    "فيدكو مصر": "Vidco Egypt",
    "شركة سعيد سعدان": "Saeed Saadan Company",
    "شركة الاندلس للتجارة": "Al Andalus Trading",
    "شركة التمساح": "Al Temsah Company",
    "شركة جمال الغمراوى": "Gamal El Ghamrawy Company",
    "الفهد استيل": "Al Fahd Steel",
}

JOBS = {
    "سائق": ("jobs.driver", "سائق", "Driver"),
    "سائق كلارك": ("jobs.forkliftDriver", "سائق كلارك", "Forklift Driver"),
    "سائق لودر": ("jobs.loaderDriver", "سائق لودر", "Loader Driver"),
    "عامل تحضيرات": ("jobs.prepWorker", "عامل تحضيرات", "Preparations Worker"),
    "فنى تحضيرات": ("jobs.prepTech", "فني تحضيرات", "Preparations Technician"),
    "مشرف تحضيرات": ("jobs.prepSupervisor", "مشرف تحضيرات", "Preparations Supervisor"),
    "رئيس قسم التحضيرات": ("jobs.prepChief", "رئيس قسم التحضيرات", "Preparations Section Head"),
    "عامل ماكينة": ("jobs.machineWorker", "عامل ماكينة", "Machine Worker"),
    "فنى ماكينة": ("jobs.machineTech", "فني ماكينة", "Machine Technician"),
    "رئيس وردية ماكينة": ("jobs.shiftLead", "رئيس وردية ماكينة", "Machine Shift Leader"),
    "مهندس ميكانيكا": ("jobs.mechEngineer", "مهندس ميكانيكا", "Mechanical Engineer"),
    "مهندس  صيانة": ("jobs.maintEngineer", "مهندس صيانة", "Maintenance Engineer"),
    "فنى صيانة ميكانيكية": ("jobs.mechTech", "فني صيانة ميكانيكية", "Mechanical Maintenance Technician"),
    "م فنى صيانة ميكانيكية": ("jobs.asstMechTech", "مساعد فني صيانة ميكانيكية", "Assistant Mechanical Technician"),
    "لحام": ("jobs.welder", "لحام", "Welder"),
    "فنى خراطة": ("jobs.turner", "فني خراطة", "Lathe Technician"),
    "رئيس قسم الصيانة": ("jobs.maintChief", "رئيس قسم الصيانة", "Maintenance Section Head"),
    "فنى كهرباء": ("jobs.elecTech", "فني كهرباء", "Electrician"),
    "م فنى كهرباء": ("jobs.asstElecTech", "مساعد فني كهرباء", "Assistant Electrician"),
    "مهندس كهرباء": ("jobs.elecEngineer", "مهندس كهرباء", "Electrical Engineer"),
    "رئيس قسم الكهرباء": ("jobs.elecChief", "رئيس قسم الكهرباء", "Electrical Section Head"),
    "فنى مقص": ("jobs.cutterTech", "فني مقص", "Cutter Technician"),
    "م فنى مقص": ("jobs.asstCutterTech", "مساعد فني مقص", "Assistant Cutter Technician"),
    "مدير جودة": ("jobs.qualityManager", "مدير جودة", "Quality Manager"),
    "مهندس جودة": ("jobs.qualityEngineer", "مهندس جودة", "Quality Engineer"),
    "عامل غلاية": ("jobs.boilerWorker", "عامل غلاية", "Boiler Worker"),
    "فنى غلاية": ("jobs.boilerTech", "فني غلاية", "Boiler Technician"),
    "م فنى غلاية": ("jobs.asstBoilerTech", "مساعد فني غلاية", "Assistant Boiler Technician"),
    "عامل تغليف": ("jobs.packer", "عامل تغليف", "Packing Worker"),
    "اخصائى سلامة وصحة مهنية": ("jobs.safetySpecialist", "أخصائي سلامة وصحة مهنية", "HSE Specialist"),
    "فنى سلامة وصحة مهنية": ("jobs.safetyTech", "فني سلامة وصحة مهنية", "HSE Technician"),
    "عامل نظافة": ("jobs.cleaner", "عامل نظافة", "Cleaner"),
    "عامل": ("jobs.laborer", "عامل", "Laborer"),
    "مدير مالى": ("jobs.financeManager", "مدير مالي", "Finance Manager"),
    "محاسب": ("jobs.accountant", "محاسب", "Accountant"),
    "مدير مخازن": ("jobs.warehouseManager", "مدير مخازن", "Warehouse Manager"),
    "امين مخزن منتج تام": ("jobs.fgStorekeeper", "أمين مخزن منتج تام", "Finished-Goods Storekeeper"),
    "امين مخزن  قطع غيار": ("jobs.spStorekeeper", "أمين مخزن قطع غيار", "Spare-Parts Storekeeper"),
    "رئيس قسم   مخزن  قطع الغيار": ("jobs.spStoreChief", "رئيس قسم مخزن قطع الغيار", "Spare-Parts Store Head"),
    "مدير مبيعات": ("jobs.salesManager", "مدير مبيعات", "Sales Manager"),
    "مدير مشتريات": ("jobs.purchasingManager", "مدير مشتريات", "Purchasing Manager"),
    "مدير موارد بشرية": ("jobs.hrManager", "مدير موارد بشرية", "HR Manager"),
    "عامل بوفية": ("jobs.buffet", "عامل بوفية", "Pantry Worker"),
}

NAME_MAP = {
    "محمد": "Mohamed", "أحمد": "Ahmed", "احمد": "Ahmed", "محمود": "Mahmoud",
    "مصطفى": "Mostafa", "مصطفي": "Mostafa", "علي": "Ali", "على": "Ali",
    "عبد": "Abdel", "الله": "Allah", "عبدالله": "Abdallah", "عبداللة": "Abdallah",
    "إبراهيم": "Ibrahim", "ابراهيم": "Ibrahim", "حسن": "Hassan", "حسين": "Hussein",
    "سعيد": "Saeed", "سيد": "Sayed", "يوسف": "Youssef", "ياسين": "Yassin",
    "يسن": "Yassin", "هشام": "Hisham", "كريم": "Karim", "اسلام": "Islam",
    "إسلام": "Islam", "وليد": "Walid", "بلال": "Bilal", "زياد": "Ziad",
    "هانى": "Hany", "هاني": "Hany", "سمير": "Samir", "صالح": "Saleh",
    "فوزي": "Fawzy", "علاء": "Alaa", "مدحت": "Medhat", "غريب": "Gharib",
    "سامي": "Samy", "سامى": "Samy", "نشأت": "Nashaat", "مؤمن": "Moamen",
    "مبروك": "Mabrouk", "رزق": "Rizk", "اشرف": "Ashraf", "أشرف": "Ashraf",
    "رضا": "Reda", "لطفي": "Lotfy", "لطفى": "Lotfy", "عمرو": "Amr",
    "جمال": "Gamal", "صلاح": "Salah", "يحيى": "Yehia", "يحى": "Yehia",
    "الحناوى": "El-Hennawy", "الحناوي": "El-Hennawy", "ماهر": "Maher",
    "مسعود": "Masoud", "حلمي": "Helmy", "عماره": "Emara", "السلام": "Salam",
    "الحميد": "Hamid", "التواب": "Tawab", "عثمان": "Osman", "وجيه": "Wagih",
    "المرضي": "Mardy", "الرزق": "Rizk", "صبري": "Sabry", "البنا": "El-Banna",
    "أيمن": "Ayman", "العزيز": "Aziz", "اللا": "Llah", "عيسي": "Eissa",
    "القاضي": "El-Kady", "الصعيدى": "El-Saidy", "الصعيدي": "El-Saidy",
    "البرقي": "El-Barqy", "عطية": "Attia", "كوس": "Kous", "الغفار": "Ghaffar",
    "ادريس": "Idris", "الزلباني": "El-Zalabany", "مختار": "Mokhtar",
    "عبده": "Abdo", "حسنين": "Hassanein", "القالع": "El-Qalea", "فريد": "Farid",
    "شتا": "Sheta", "أمين": "Amin", "عاشور": "Ashour", "حمدي": "Hamdy",
    "حمزة": "Hamza", "النبي": "Nabi", "الحضري": "El-Hadary", "عادل": "Adel",
    "كشك": "Keshek", "السعيد": "El-Saeed", "السيد": "El-Sayed", "مرسي": "Morsy",
    "القشيط": "El-Qosheit", "عيد": "Eid", "الحبال": "El-Hebal", "الدكروري": "El-Dakroury",
    "السطوحى": "El-Sattouhy", "عبيد": "Ebeid", "صبرة": "Sabra", "الظاهر": "Zaher",
    "المنسى": "El-Mensy", "الحليم": "Halim", "معروف": "Maarouf", "صباح": "Sabah",
    "فايز": "Fayez", "عطا": "Atta", "حسنى": "Hosny", "السميع": "Sami",
    "امين": "Amin", "الخالق": "Khalek", "سعد": "Saad", "عميرة": "Omaira",
    "نجيب": "Naguib", "فهمى": "Fahmy", "عمر": "Omar", "الغراب": "El-Ghorab",
    "تعيلب": "Taalab", "المسدى": "El-Masdy", "شلتوت": "Shaltout", "الجداوى": "El-Gadawy",
    "اسماعيل": "Ismail", "إسماعيل": "Ismail", "العباسى": "El-Abbasy", "مغربي": "Maghraby",
    "زيد": "Zeid", "العليم": "Aleem", "العظيم": "Azim", "نعيم": "Naeem",
    "المقصود": "Maqsoud", "موسى": "Moussa", "الجواد": "Gawad", "العمروس": "El-Amroos",
    "بكر": "Bakr", "جمعة": "Gomaa", "حامد": "Hamed", "الدغيدى": "El-Degheidy",
    "سمعان": "Samaan", "حنا": "Hanna", "عوض": "Awad", "قسطندى": "Kostandy",
    "جرجس": "Gerges", "القط": "El-Qott", "حمودة": "Hamouda", "شوقى": "Shawky",
    "شحاتة": "Shehata", "نبوى": "Nabawy", "شامل": "Shamel", "الهادى": "Hady",
    "تر ك": "Tork", "ترك": "Tork", "يسرى": "Yousry", "الصادق": "Sadek",
    "خليل": "Khalil", "شوقي": "Shawky", "عنتر": "Antar", "العشرى": "El-Ashry",
    "نوح": "Nouh", "شوكت": "Shawkat", "الحسينى": "El-Husseiny", "الشريف": "El-Sherif",
    "المحسن": "Mohsen", "صابر": "Saber", "القزاز": "El-Qazzaz", "خالد": "Khaled",
    "رضوان": "Radwan", "بشير": "Bashir", "فليفل": "Fleifel", "النحاس": "El-Nahhas",
    "السنوسى": "El-Senousy", "عيسوى": "Eissawy", "كيروان": "Kerwan", "هيكل": "Heikal",
    "الامبابى": "El-Embaby", "صبحى": "Sobhy", "عثمان": "Osman", "ابو": "Abu",
    "المجد": "Magd", "حشاد": "Hashad", "السادات": "El-Sadat", "السيسى": "El-Sisi",
    "خضر": "Khedr", "طه": "Taha", "طة": "Taha", "الرازق": "Razek",
    "الرسول": "Rasoul", "اللاه": "Llah",
}

LETTERS = {
    "ا": "a", "أ": "a", "إ": "i", "آ": "a", "ب": "b", "ت": "t", "ث": "th",
    "ج": "g", "ح": "h", "خ": "kh", "د": "d", "ذ": "z", "ر": "r", "ز": "z",
    "س": "s", "ش": "sh", "ص": "s", "ض": "d", "ط": "t", "ظ": "z", "ع": "a",
    "غ": "gh", "ف": "f", "ق": "q", "ك": "k", "ل": "l", "م": "m", "ن": "n",
    "ه": "h", "و": "w", "ي": "y", "ى": "a", "ة": "a", "ء": "", "ئ": "y",
    "ؤ": "o",
}

FOREIGN_AR = {
    "Weifang idol technology co ltd": "ويفانغ آيدول للتكنولوجيا",
    "BINET SUL LIRI S.P.A.": "بينيه سول ليري",
    "J.M.Voith SE & Co. KG": "فويث",
    "VILLFORTH Siebtechnik GmbH": "فيلفورت لتقنية الغرابيل",
    "Greatland pulp and paper technology": "غريت لاند لتقنية اللب والورق",
    "WEIFANG GREATLAND CHEMICALS CO., LIMITED": "ويفانغ غريت لاند للكيماويات",
    "Vipa laussane sa": "فيبا لوزان",
    "MARE DYNAMICS SRL": "ماري دايناميكس",
    "WEINGRILL S.R.L.": "وينغريل",
    "Paul Wegner GmbH & Co. KG": "بول فيغنر",
    "AK-PA TEKSTİL İHRACAT PAZARLAMA A.Ş.": "آك با للتسويق والتصدير",
    "PowerMade S.r.l.": "باور ميد",
    "overmade srl": "أوفر ميد",
    "Ing. Christian Gruber": "المهندس كريستيان غروبر",
    "EMCO DYESTUFF PVT LTD.": "إمكو للأصباغ",
    "KEMIND SRL": "كيميند",
    "SHALIMAR WIRES INDUSTRIES LIMITED": "شاليمار للأسلاك",
    "BASF FZE": "باسف",
    "SABIC": "سابك",
}


def js(value: object) -> str:
    if value is None:
        return "undefined"
    text = str(value)
    return "'" + text.replace("\\", "\\\\").replace("'", "\\'") + "'"


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def translate_account(name: str) -> str:
    out = name
    for en, ar in sorted(PHRASES, key=lambda pair: -len(pair[0])):
        out = re.sub(re.escape(en), ar, out, flags=re.I)
    parts = []
    for token in re.split(r"(\s+|[–\-/,_])", out):
        if token in WORDS:
            if WORDS[token]:
                parts.append(WORDS[token])
        else:
            parts.append(token)
    ar = re.sub(r"\s+", " ", "".join(parts)).strip(" -–/")
    return ar or name


def transliterate(name: str) -> str:
    bits = []
    for part in name.split():
        if part in NAME_MAP:
            bits.append(NAME_MAP[part])
            continue
        mapped = "".join(LETTERS.get(ch, ch) for ch in part)
        bits.append(mapped[:1].upper() + mapped[1:] if mapped else part)
    return " ".join(bits)


DEPT_KEY = {
    "الحركة": "administrations.fleet",
    "الانتاج": "administrations.production",
    "الصيانة": "administrations.maintenance",
    "الصيانة الكهربائية": "administrations.electrical",
    "الادارة المالية": "administrations.finance",
    "الجودة": "administrations.quality",
    "السلامة والصحة المهنية": "administrations.safety",
}

SECTION_KEY = {
    ("الحركة", "الحركة"): "sections.fleet",
    ("الانتاج", "التحضيرات"): "sections.prep",
    ("الانتاج", "الجودة"): "sections.prodQuality",
    ("الانتاج", "الماكينة"): "sections.machine",
    ("الصيانة", "الصيانة الميكانيكية"): "sections.mechanical",
    ("الصيانة الكهربائية", "الكهرباء"): "sections.electrical",
    ("الادارة المالية", "الادارة"): "sections.office",
    ("الادارة المالية", "الحسابات"): "sections.accounts",
    ("الادارة المالية", "المبيعات"): "sections.sales",
    ("الادارة المالية", "المخازن"): "sections.warehouse",
    ("الادارة المالية", "المشتريات"): "sections.purchasing",
    ("الادارة المالية", "الموارد البشرية"): "sections.hr",
    ("الجودة", "الجودة"): "sections.quality",
    ("السلامة والصحة المهنية", "السلامة والصحة المهنية"): "sections.hse",
}


def dept_of(dept: str, _section: str) -> str:
    return DEPT_KEY.get(clean(dept), "departments.finance")


def section_of(dept: str, section: str) -> str:
    return SECTION_KEY.get((clean(dept), clean(section)), "")


def role_of(job_key: str, dept_key: str) -> str:
    if job_key in {"jobs.financeManager", "jobs.accountant"}:
        return "finance"
    if job_key in {"jobs.warehouseManager", "jobs.fgStorekeeper", "jobs.spStorekeeper", "jobs.spStoreChief"}:
        return "store"
    if job_key == "jobs.hrManager":
        return "admin"
    return "operator"


def nature_of(acc_type: str, name: str) -> str:
    if re.search(r"purchase", name, re.I) and "return" not in name.lower():
        return "debit"
    if re.search(r"sales return", name, re.I):
        return "debit"
    if acc_type in {"Asset", "Expense"}:
        return "debit"
    return "credit"


def parent_of(code: str, codes: set[str]) -> str | None:
    hits = [other for other in codes if other != code and code.startswith(other)]
    return max(hits, key=len) if hits else None


def load_sheet(path: Path, sheet: str | None = None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")
    lines = text.count("\n") + 1
    if lines > 300:
        raise SystemExit(f"{path} has {lines} lines")


def gen_accounts() -> None:
    rows = load_sheet(DOWNLOADS / "Chart_of_Accounts.xlsx", "Chart_of_Accounts")
    header = next(i for i, row in enumerate(rows) if row and row[0] == "Account Code *")
    raw: list[dict] = []
    seen: Counter[str] = Counter()
    for row in rows[header + 1 :]:
        if row[0] is None:
            continue
        code = clean(row[0])
        seen[code] += 1
        if seen[code] > 1:
            code = f"{code}-{seen[code]}"
        raw.append(
            {
                "code": code,
                "en": clean(row[1]),
                "type": clean(row[3]),
                "level": int(row[4] or 1),
                "ccy": clean(row[6]) or "EGP",
            }
        )
    codes = {row["code"] for row in raw}
    children = Counter(parent_of(row["code"], codes) or "" for row in raw)
    chunks: list[list[str]] = [[], [], []]
    for index, row in enumerate(raw):
        parent = parent_of(row["code"], codes)
        line = (
            f"  acc({js(row['code'])}, {js(translate_account(row['en']))}, {js(row['en'])}, "
            f"{js(parent) if parent else 'undefined'}, {row['level']}, "
            f"{js(nature_of(row['type'], row['en']))}, {js(row['ccy'])}, "
            f"{'true' if children[row['code']] == 0 else 'false'}),"
        )
        chunks[min(index * 3 // len(raw), 2)].append(line)
    util = """import { Account, AccountNature } from '../../../core/models/finance.models';

export function acc(
  code: string,
  name: string,
  nameEn: string,
  parentCode: string | undefined,
  level: number,
  nature: AccountNature,
  currency: string,
  isPostable: boolean,
): Account {
  return {
    code,
    name,
    name_en: nameEn,
    parentCode,
    level,
    nature,
    currency,
    isPostable,
    costCenterRequired: false,
  };
}
"""
    write(SEED / "accounts.util.ts", util)
    for index, lines in enumerate(chunks, 1):
        body = "\n".join(
            [
                "import { Account } from '../../../core/models/finance.models';",
                "import { acc } from './accounts.util';",
                "",
                f"export const SEED_ACCOUNTS_{index}: Account[] = [",
                *lines,
                "];",
                "",
            ]
        )
        write(SEED / f"accounts-{index}.ts", body)


def gen_customers() -> None:
    rows = load_sheet(DOWNLOADS / "Customers.xlsx", "Customers")
    header = next(i for i, row in enumerate(rows) if row and row[1] == "Customer Name *")
    lines = [
        "import { Customer } from '../../../core/models/sales.models';",
        "",
        "export const SEED_CUSTOMERS: Customer[] = [",
    ]
    n = 0
    for row in rows[header + 1 :]:
        en = clean(row[1])
        if not en:
            continue
        n += 1
        ar, city_ar = CUS_AR.get(en, (en, clean(row[7])))
        city_en = clean(row[7]) or city_ar
        ccy = clean(row[13]).upper() or "EGP"
        if ccy not in {"EGP", "USD", "EUR"}:
            ccy = "USD" if "export" in clean(row[2]).lower() else "EGP"
        lines.append(
            "  { "
            f"code: 'CUS-{n:03d}', name: {js(ar)}, name_en: {js(en)}, "
            f"region: {js(city_ar)}, region_en: {js(city_en)}, currency: {js(ccy)}, balance: 0 "
            "},"
        )
    lines.append("];")
    write(SEED / "customers.seed.ts", "\n".join(lines) + "\n")


def currency_of(raw: object) -> str:
    text = clean(raw).upper()
    if "USD" in text or "DOLLAR" in text:
        return "USD"
    if "EUR" in text or "EURO" in text:
        return "EUR"
    return "EGP"


def gen_suppliers() -> None:
    lines = [
        "import { Supplier } from '../../../core/models/purchasing.models';",
        "",
        "export const SEED_SUPPLIERS: Supplier[] = [",
    ]
    n = 0

    def add(name_ar: str, name_en: str, ccy: str) -> None:
        nonlocal n
        n += 1
        lines.append(
            "  { "
            f"code: 'SUP-{n:03d}', name: {js(name_ar)}, name_en: {js(name_en)}, "
            f"currency: {js(ccy)}, balance: 0, onTimeDeliveryPercent: 90 "
            "},"
        )

    local_rows = load_sheet(DOWNLOADS / "AlHennawy_ERP_Master_Data_Collection.xlsx", "Suppliers")
    header = next(i for i, row in enumerate(local_rows) if row and "Supplier Name" in str(row[1] or ""))
    for row in local_rows[header + 1 :]:
        name = clean(row[1])
        if not name:
            continue
        add(name, SUP_EN.get(name, name), currency_of(row[13] if len(row) > 13 else "EGP"))

    foreign_rows = load_sheet(DOWNLOADS / "SUPPLIERS.xlsx", "Suppliers")
    header = next(i for i, row in enumerate(foreign_rows) if row and row[1] == "Supplier Name *")
    for row in foreign_rows[header + 1 :]:
        en = clean(row[1]).rstrip(",")
        if not en:
            continue
        add(FOREIGN_AR.get(en, en), en, currency_of(row[13] if len(row) > 13 else "USD"))
    lines.append("];")
    write(SEED / "suppliers.seed.ts", "\n".join(lines) + "\n")


def gen_employees() -> None:
    rows = load_sheet(DOWNLOADS / "اسماء العاملين بشركة الحناوى.xlsx", "الحضور")
    lines = [
        "import { Employee } from '../../../core/models/hr.models';",
        "",
        "const IMG = 'assets/branding/alhennawy-logo.png';",
        "",
        "function emp(",
        "  id: string,",
        "  code: string,",
        "  name: string,",
        "  nameEn: string,",
        "  departmentKey: string,",
        "  sectionKey: string,",
        "  jobTitleKey: string,",
        "  roleId: string,",
        "): Employee {",
        "  return {",
        "    id, code, name, name_en: nameEn, email: `${code.toLowerCase()}@alhennawy.net`,",
        "    password: 'emp123', departmentKey, sectionKey, jobTitleKey, hireDate: '2020-01-01',",
        "    status: 'active', leaveBalanceDays: 21, photoUrl: IMG, drugTestImageUrl: IMG,",
        "    roleId, workStart: '08:00', workEnd: '16:00',",
        "  };",
        "}",
        "",
        "export const SEED_EMPLOYEES: Employee[] = [",
    ]
    seen: set[str] = set()
    for row in rows[2:]:
        code = clean(row[1])
        name = clean(row[2])
        if not name or not code or "%" in code or "يومية" in code:
            continue
        if code in seen:
            continue
        seen.add(code)
        job = clean(row[5])
        job_key, _, _ = JOBS.get(job, ("jobs.laborer", "عامل", "Laborer"))
        dept = dept_of(clean(row[3]), clean(row[4]))
        section = section_of(clean(row[3]), clean(row[4]))
        lines.append(
            f"  emp({js('e-' + code)}, {js(code)}, {js(name)}, {js(transliterate(name))}, "
            f"{js(dept)}, {js(section)}, {js(job_key)}, {js(role_of(job_key, dept))}),"
        )
    lines.append("];")
    write(SEED / "employees.seed.ts", "\n".join(lines) + "\n")


def gen_jobs_lookups() -> None:
    lines = [
        "import { LookupValue } from '../../../core/models/system.models';",
        "",
        "export const SEED_JOB_LOOKUPS: LookupValue[] = [",
    ]
    used = set()
    for key, ar, en in JOBS.values():
        if key in used:
            continue
        used.add(key)
        lines.append(
            f"  {{ id: {js('lv-job-' + key.split('.')[-1])}, group: 'jobTitles', "
            f"value: {js(key)}, labelAr: {js(ar)}, labelEn: {js(en)} }},"
        )
    lines.append("];")
    write(SEED / "job-lookups.seed.ts", "\n".join(lines) + "\n")


if __name__ == "__main__":
    gen_accounts()
    gen_customers()
    gen_suppliers()
    gen_employees()
    gen_jobs_lookups()
    print("seed files written under", SEED)
